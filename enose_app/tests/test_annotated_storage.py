# -*- coding: utf-8 -*-
"""验证原始帧与异步推理状态合并后，CSV/XLSX保存结果一致。"""

import csv
import sys
import tempfile
import unittest
from pathlib import Path


APP_ROOT = Path(__file__).resolve().parents[1]
PUBLIC_ROOT = APP_ROOT / "public"
if str(PUBLIC_ROOT) not in sys.path:
    sys.path.insert(0, str(PUBLIC_ROOT))

from robot_app.storage import (  # noqa: E402
    AsyncFrameRecorder,
    COLUMNS,
    InferenceFrameJoiner,
    PHASE_RESULTS,
)


def make_frame(sequence, timestamp_ms=None):
    """生成一帧结构完整、数值容易核对的测试数据。"""
    sequence = int(sequence)
    return {
        "timestamp_ms": int(timestamp_ms if timestamp_ms is not None else 1000 + sequence),
        "device_id": "回流1",
        "sequence": sequence,
        "resistance": [10.0 * sequence + index for index in range(1, 17)],
        "temperature_c": 25.0 + sequence / 10.0,
        "humidity_rh": 40.0 + sequence / 10.0,
    }


class MemoryRecorder:
    """只记录调用参数，隔离磁盘写入以便测试关联规则。"""

    def __init__(self):
        self.rows = []

    def record(self, frame, recognition_result=None):
        self.rows.append((dict(frame), recognition_result))


class InferenceFrameJoinerTests(unittest.TestCase):
    def test_all_phases_are_saved_as_expected_chinese_status(self):
        recorder = MemoryRecorder()
        joiner = InferenceFrameJoiner(recorder, max_pending=10)

        for sequence, phase in enumerate(PHASE_RESULTS):
            joiner.add_frame(make_frame(sequence))
            self.assertTrue(
                joiner.handle_status(
                    {"device_id": "回流1", "sequence": sequence, "phase": phase}
                )
            )

        self.assertEqual(
            [result for _frame, result in recorder.rows],
            [PHASE_RESULTS[phase] for phase in PHASE_RESULTS],
        )

    def test_final_label_has_priority_over_status(self):
        recorder = MemoryRecorder()
        joiner = InferenceFrameJoiner(recorder)
        joiner.add_frame(make_frame(8))

        self.assertTrue(
            joiner.handle_result(
                {
                    "device_id": "回流1",
                    "confirm_sequence": 8,
                    "label": "GETARIA",
                }
            )
        )
        self.assertTrue(
            joiner.handle_status(
                {"device_id": "回流1", "sequence": 8, "phase": "recovering"}
            )
        )

        self.assertEqual(len(recorder.rows), 1)
        self.assertEqual(recorder.rows[0][1], "GETARIA")

    def test_recovery_event_has_priority_and_next_frame_returns_to_idle(self):
        recorder = MemoryRecorder()
        joiner = InferenceFrameJoiner(recorder)
        recovery_frame = make_frame(20, timestamp_ms=20000)
        idle_frame = make_frame(21, timestamp_ms=21000)
        joiner.add_frame(recovery_frame)
        joiner.add_frame(idle_frame)

        self.assertTrue(
            joiner.handle_recovery(
                {
                    "device_id": "回流1",
                    "source_timestamp_ms": 20000,
                    "status": "complete",
                }
            )
        )
        joiner.handle_status(
            {"device_id": "回流1", "sequence": 20, "phase": "idle"}
        )
        joiner.handle_status(
            {"device_id": "回流1", "sequence": 21, "phase": "idle"}
        )

        self.assertEqual(
            [result for _frame, result in recorder.rows],
            ["恢复完成", "等待信号"],
        )

    def test_timeout_and_error_are_saved(self):
        recorder = MemoryRecorder()
        joiner = InferenceFrameJoiner(recorder)
        joiner.add_frame(make_frame(30, timestamp_ms=30000))
        joiner.add_frame(make_frame(31, timestamp_ms=31000))

        joiner.handle_recovery(
            {
                "device_id": "回流1",
                "source_timestamp_ms": 30000,
                "status": "timeout",
            }
        )
        joiner.handle_status(
            {"device_id": "回流1", "sequence": 30, "phase": "recovering"}
        )
        self.assertTrue(joiner.handle_error({"source_sequence": 31}))

        self.assertEqual(
            [result for _frame, result in recorder.rows],
            ["恢复超时", "推理错误"],
        )

    def test_overflow_and_shutdown_keep_every_raw_frame(self):
        recorder = MemoryRecorder()
        warnings = []
        joiner = InferenceFrameJoiner(recorder, max_pending=2, on_warning=warnings.append)

        joiner.add_frame(make_frame(0))
        joiner.add_frame(make_frame(1))
        joiner.add_frame(make_frame(2))
        self.assertEqual(len(recorder.rows), 1)
        self.assertEqual(recorder.rows[0][1], "未收到推理状态")
        self.assertEqual(len(warnings), 1)

        self.assertEqual(joiner.flush_unresolved(), 2)
        self.assertEqual([row[0]["sequence"] for row in recorder.rows], [0, 1, 2])
        self.assertEqual(
            [result for _frame, result in recorder.rows],
            ["未收到推理状态"] * 3,
        )


class FileSchemaTests(unittest.TestCase):
    def _read_parts(self, output_dir, output_format):
        rows = []
        paths = sorted(output_dir.glob("*.%s" % output_format))
        self.assertEqual(len(paths), 2)
        for path in paths:
            if output_format == "csv":
                with path.open("r", encoding="utf-8-sig", newline="") as stream:
                    values = list(csv.reader(stream))
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(path, read_only=True, data_only=True)
                worksheet = workbook["raw_data"]
                values = [list(row) for row in worksheet.iter_rows(values_only=True)]
                workbook.close()
            self.assertEqual(list(values[0]), list(COLUMNS))
            rows.extend(values[1:])
        return rows

    def test_csv_and_xlsx_have_the_same_new_last_column_and_no_lost_rows(self):
        for output_format in ("csv", "xlsx"):
            with self.subTest(output_format=output_format), tempfile.TemporaryDirectory() as temp:
                output_dir = Path(temp)
                recorder = AsyncFrameRecorder(
                    output_dir,
                    output_format=output_format,
                    rotate_rows=2,
                    xlsx_flush_every=1,
                )
                expected_results = ["预热中", "GETARIA", "恢复完成"]
                for sequence, result in enumerate(expected_results):
                    recorder.record(make_frame(sequence), result)
                recorder.close()

                self.assertEqual(recorder.rows_written, 3)
                rows = self._read_parts(output_dir, output_format)
                self.assertEqual(len(rows), 3)
                self.assertEqual(len(COLUMNS), 22)
                self.assertEqual(COLUMNS[-1], "Recognition_Result")
                self.assertEqual([int(row[2]) for row in rows], [0, 1, 2])
                self.assertEqual([row[-1] for row in rows], expected_results)
                self.assertEqual(float(rows[2][3]), make_frame(2)["resistance"][0])
                self.assertEqual(float(rows[2][19]), make_frame(2)["temperature_c"])
                self.assertEqual(float(rows[2][20]), make_frame(2)["humidity_rh"])


if __name__ == "__main__":
    unittest.main()
